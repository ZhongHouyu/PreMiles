"""生成**国赛格式的 Word 论文**，公式为 Word 原生公式对象（OMML，MathType 可一键转换）。

内容与 `build_cumcm_paper.py`（PDF 版）同源：同一次 `split_source()` / `add_citations()`，
摘要、参考文献也复用同一份常量，避免两份产物内容漂移。

版式（同国赛规范）
------------------
* A4、四边 2.5 cm；正文 小四号宋体 + Times New Roman、1.5 倍行距
* 论文题目 三号黑体居中；一级标题 四号黑体居中；二/三级标题 小四号黑体左对齐
* 图表五号；图题在下、表题在上，连续编号
* 页码自摘要页起，页脚居中；承诺书页与编号专用页不编码
* **公式全部为 OMML 公式对象**：Word 里可直接编辑；装 MathType 后用
  「Convert Equations → Office Math」一键转成 MathType 公式

用法
----
    python build_cumcm_docx.py
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_cumcm_paper as cumcm
import math_omml as mo

HERE = Path(__file__).resolve().parent
OUT = HERE / "out"
DOCX_NAME = "论文_整车能耗归因与降耗方向（数学建模国赛版）.docx"

SONG = "宋体"
HEI = "黑体"
LATIN = "Times New Roman"
MONO = "Consolas"

# 公式标记：私有区字符，Word 转换脚本据此识别并生成原生公式，转换后删除
MK_IN_S, MK_IN_E = "\ue000", "\ue001"      # 行内公式
MK_DI_S, MK_DI_E = "\ue002", "\ue003"      # 展示式公式

# 数学斜体 / 希腊（行内识别用）
MATH_ITALIC = re.compile(r"[\U0001D400-\U0001D7FF\u210E]")
GREEK_CH = "ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩαβγδεζηθικλμνξοπρστυφχψψω"
STRONG_SYMBOLS = "=∫∑∝≤≥"
CJK = re.compile(r"[\u4e00-\u9fff]")


# ------------------------------------------------------------------ 基础工具
def set_run(run, size=12.0, ea=SONG, latin=LATIN, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), ea)


def para(doc, text="", size=12.0, ea=SONG, latin=LATIN, bold=False,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.5, indent_first=False,
         space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent_first:
        pf.first_line_indent = Pt(size * 2)
    if text:
        set_run(p.add_run(text), size=size, ea=ea, latin=latin, bold=bold)
    return p


def add_omml(paragraph, xml: str) -> None:
    paragraph._p.append(parse_xml(xml))


def add_page_field(paragraph, size=10.5):
    run = paragraph.add_run()
    set_run(run, size=size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def setup_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.5))


def strip_html(html: str) -> list[str]:
    """把 HTML 片段转成纯文本段落列表。"""
    text = re.sub(r"<br\s*/?>", "\n", html)
    text = re.sub(r"</p>", "\n", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = (text.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&"))
    return [ln.strip() for ln in text.split("\n") if ln.strip()]


# ------------------------------------------------------------------ 行内渲染
# 行内数学 token：以数学斜体字母 / 希腊字母开头（也允许 "(" 紧跟其后），
# 用**带括号配平的手写扫描**向后扩展，避免把外层的右括号吞进公式
# （例如 "(1 + δ)" 只应把 δ 当公式；"F)²" 这种畸形 token 必须避免）。
_MATH_START = set("".join(chr(c) for c in range(0x1D400, 0x1D800))) | {"\u210e"} | set("ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩαβγδεζηθικλμνξοπρστυφχψω")
_TAIL_CHARS = set("_^+-−·/×²³⁻¹⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉"
                  "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789") \
    | _MATH_START

OTHER_RE = re.compile(
    r"(\*\*[^*\n]+\*\*)"          # 粗体
    r"|(`[^`\n]+`)"               # 行内代码
    r"|(\[[^\]]*\]\([^)\s]+\))"   # 链接
    r"|(!\[[^\]]*\]\([^)\s]+\))"  # 图片
)


def _is_math_start(text: str, i: int) -> bool:
    ch = text[i]
    if ch in _MATH_START:
        return True
    return ch == "(" and i + 1 < len(text) and text[i + 1] in _MATH_START


def _extend_math(text: str, i: int) -> int:
    """从 i 扩展出一个行内数学 token，返回结束位置（括号配平，不吞外层右括号）。"""
    depth = 0
    j = i
    if text[j] == "(":
        depth = 1
        j += 1
    while j < len(text):
        ch = text[j]
        if ch == "(":
            depth += 1
            j += 1
            continue
        if ch == ")":
            if depth == 0:
                break
            depth -= 1
            j += 1
            continue
        if ch == "_" and j + 1 < len(text) and "\u4e00" <= text[j + 1] <= "\u9fff":
            k = j + 1
            while k < len(text) and "\u4e00" <= text[k] <= "\u9fff":
                k += 1
            j = k
            continue
        if ch in _TAIL_CHARS:
            j += 1
            continue
        break
    return j


def _render_plain(p, text: str, size: float, bold: bool = False) -> None:
    """普通文本片段：把其中的行内公式转成 OMML，其余当正文。"""
    pos = 0
    i = 0
    while i < len(text):
        if not _is_math_start(text, i):
            i += 1
            continue
        end = _extend_math(text, i)
        if end <= i:
            i += 1
            continue
        if i > pos:
            set_run(p.add_run(text[pos:i]), size=size, bold=bold)
        add_omml(p, mo.omml_inline_xml(text[i:end]))
        pos = end
        i = end
    if pos < len(text):
        set_run(p.add_run(text[pos:]), size=size, bold=bold)


def render_inline(p, text: str, size: float):
    """把一行 Markdown 文本写进段落，含行内公式（OMML）。"""
    pos = 0
    for m in OTHER_RE.finditer(text):
        if m.start() > pos:
            _render_plain(p, text[pos:m.start()], size)
        token = m.group(0)
        if m.group(1):                                   # 粗体：内部同样要处理行内公式
            _render_plain(p, token[2:-2], size, bold=True)
        elif m.group(2):                                 # 行内代码
            set_run(p.add_run(token[1:-1]), size=size - 0.5, latin=MONO)
        elif m.group(3):                                 # 链接
            label = re.match(r"\[([^\]]*)\]", token)
            set_run(p.add_run(label.group(1) if label else token), size=size)
        else:                                            # 图片（此处不处理）
            set_run(p.add_run(re.match(r"!\[([^\]]*)\]", token).group(1)), size=size)
        pos = m.end()
    if pos < len(text):
        _render_plain(p, text[pos:], size)


def is_display_formula(text: str) -> bool:
    """判断整段是否为展示式（居中公式）。"""
    if "**" in text:
        # 带粗体标记的行是“**标签**：说明/公式”的散文，必须走行内渲染，
        # 否则 ** 会被当成公式内容，在 Word 里显示成字面星号。
        return False
    plain = re.sub(r"[*`]", "", text).strip()
    if not plain or len(plain) > 220:
        return False
    cjk = len(CJK.findall(plain))
    has_eq = any(s in plain for s in STRONG_SYMBOLS)
    has_mul = "·" in plain and len(MATH_ITALIC.findall(plain)) >= 1
    if cjk > 8:
        return False
    return bool(has_eq or has_mul) and len(MATH_ITALIC.findall(plain)) + len(re.findall(r"[=∫∑·]", plain)) >= 2


_NOTE_CJK = re.compile(r"[\u4e00-\u9fff]")
_EQNO = re.compile(r"^\s*[0-9]{1,3}[a-z]?\s*$")


def _is_note(content: str) -> bool:
    """判断一对括号是不是"公式末尾的注释"。

    只有两种情况算注释：公式编号（如 (6)）与含中文的说明（如 （𝐾 = 5 折，shuffle））。
    纯数学的括号组是公式本体的一部分（例如 `abs(𝑟 / 𝑦)`、`sd(𝑟)`、`(1 / 𝑛)`），
    一旦误判成注释，公式里就会只剩一个孤零零的 `abs`。
    """
    inner = content.strip()[1:-1].strip()
    if not inner:
        return False
    if _EQNO.match(inner):
        return True
    return bool(_NOTE_CJK.search(inner))


def split_trailing_note(text: str) -> tuple[str, str]:
    """把公式末尾的注释（编号 / 中文说明）拆出来，避免它们进入公式对象。"""
    formula, notes = text, []
    # 展示式末尾可能连着两段注释：编号 (6) 与中文说明（…），从右往左逐段剥离
    for _ in range(2):
        m = re.search(r"[（(][^（）()]*[）)]\s*$", formula)
        if not m or not _is_note(m.group(0)):
            break
        notes.insert(0, m.group(0).strip())
        formula = formula[:m.start()].strip()
    return formula, "　".join(notes)


def clean_formula_text(text: str) -> str:
    """去掉公式文本里的 Markdown/HTML 残留（`<br>`、粗体/斜体标记等）。

    这些标记如果混进 OMML，Word 会把它们当字面字符排版出来（例如公式里出现 `<br>` 或 `**`）。
    """
    t = re.sub(r"<br\s*/?>", " ", text)      # 公式里的换行标记 → 空格（多行等式拆成多行展示式）
    t = re.sub(r"</?[A-Za-z][^>]*>", "", t)  # 其它 HTML 标签
    t = t.replace("**", "").replace("`", "")
    t = t.replace("*", "")                   # 单个 * 在本论文里不是乘号（乘号写作 · 或 ×）
    return re.sub(r"\s{2,}", " ", t).strip()


# ------------------------------------------------------------------ 文档构建
def build() -> Path:
    OUT.mkdir(exist_ok=True)
    src = cumcm.split_source()
    body_md = cumcm.add_citations(src["body"] + "\n\n" + src["appendix"])
    abstract = strip_html(cumcm.ABSTRACT)
    references = strip_html(cumcm.REFERENCES)

    doc = Document()
    setup_page(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = LATIN
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)

    # ---------------- 第 1 页：承诺书 ----------------
    para(doc, "承　诺　书", size=16, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.5, space_before=36, space_after=24)
    for line in strip_html(cumcm.COMMITMENT):
        para(doc, line, size=12, spacing=1.5, space_after=10)
    for line in ["我们参赛选择的题号是（从 A/B/C/D/E 中选择一项填写）：＿＿＿＿＿＿＿＿＿＿",
                 "我们的参赛报名号为（如果赛区设置报名号的话）：＿＿＿＿＿＿＿＿＿＿",
                 "所属学校（请填写完整的全名）：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿",
                 "参赛队员（打印并签名）：1. ＿＿＿＿＿＿　2. ＿＿＿＿＿＿　3. ＿＿＿＿＿＿",
                 "指导教师或指导教师组负责人（打印并签名）：＿＿＿＿＿＿＿＿＿＿"]:
        p = para(doc, line, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=2.0, space_after=8)
    para(doc, "日期：＿＿＿＿年＿＿月＿＿日", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT,
         spacing=2.0, space_before=12)

    # ---------------- 第 2 页：编号专用页 ----------------
    doc.add_page_break()
    para(doc, "编　号　专　用　页", size=14, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_before=36, space_after=36)
    para(doc, "赛区评阅编号（由赛区组委会评阅前进行编号）：", size=12, spacing=2.0)
    para(doc, "", size=12, spacing=2.0, space_after=24)
    para(doc, "全国评阅编号（由全国组委会评阅前进行编号）：", size=12, spacing=2.0)

    # ---------------- 正文（新节：页码从摘要页 = 1）----------------
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_page(section)
    sectPr = section._sectPr
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    sectPr.append(pgNumType)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(fp)

    # 摘要页
    para(doc, src["title"], size=16, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.4, space_after=4)
    para(doc, src["sub"], size=12, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.4, space_after=14)
    para(doc, "摘　要", size=14, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.4, space_after=10)
    for chunk in abstract:
        p = para(doc, "", size=12, spacing=1.5, space_after=6)
        render_inline(p, chunk, 12)

    doc.add_page_break()

    # ---------------- 正文 ----------------
    fig_no = tbl_no = 0
    last_head = ""          # 最近一个标题，用作表题
    lines = body_md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        # 代码块
        if re.match(r"^\s*```", line):
            i += 1
            buf = []
            while i < len(lines) and not re.match(r"^\s*```", lines[i]):
                buf.append(lines[i])
                i += 1
            i += 1
            for b in buf:
                para(doc, b, size=9, latin=MONO, align=WD_ALIGN_PARAGRAPH.LEFT,
                     spacing=1.15, space_after=0)
            para(doc, "", size=6, space_after=6)
            continue
        # 标题
        m = re.match(r"^(#{2,4})\s+(.*?)\s*#*\s*$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            last_head = re.sub(r"^[0-9.、\s]+", "", text)
            if level == 2:
                para(doc, text, size=14, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
                     spacing=1.5, space_before=14, space_after=8)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(6)
                render_inline(p, text, 12)
                for r in p.runs:
                    set_run(r, size=12, ea=HEI)
            i += 1
            continue
        # 表格
        if "|" in line and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            tbl_no += 1
            subject = last_head[:34]
            para(doc, f"表 {tbl_no}　{subject}", size=10.5,
                 align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.2, space_after=2)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for k, cell_text in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.text = ""
                render_inline(cell.paragraphs[0], cell_text, 10.5)
                for r in cell.paragraphs[0].runs:
                    set_run(r, size=10.5, ea=HEI)
            for row in rows:
                cells = table.add_row().cells
                for k in range(len(header)):
                    txt = row[k] if k < len(row) else ""
                    cells[k].text = ""
                    render_inline(cells[k].paragraphs[0], txt, 10.5)
            para(doc, "", size=6, space_after=6)
            continue
        # 引用块
        if re.match(r"^\s*>", line):
            buf = []
            while i < len(lines) and re.match(r"^\s*>", lines[i]):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for b in buf:
                if not b.strip():
                    continue
                p = para(doc, "", size=11, spacing=1.4, space_after=4)
                p.paragraph_format.left_indent = Cm(0.6)
                render_inline(p, b, 11)
            continue
        # 列表
        m = re.match(r"^\s*([-*+]|\d{1,3}[.)])\s+(.*)$", line)
        if m:
            marker = "·" if not m.group(1)[0].isdigit() else m.group(1)
            text = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            set_run(p.add_run(f"{marker} "), size=12)
            render_inline(p, text, 12)
            i += 1
            continue
        # 分隔线
        if re.match(r"^\s*(?:[-*_]\s*){3,}$", line):
            i += 1
            continue
        # 图片
        m = re.match(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$", line.strip())
        if m:
            alt, rel = m.group(1), m.group(2)
            img = (HERE / rel).resolve()
            if img.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                p.add_run().add_picture(str(img), width=Cm(14.5))
                fig_no += 1
                para(doc, f"图 {fig_no}　{alt}", size=10.5,
                     align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.2, space_after=8)
            i += 1
            continue
        # 展示式公式
        if is_display_formula(line):
            formula, note = split_trailing_note(clean_formula_text(line.strip()))
            if not formula:
                i += 1
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_omml(p, mo.omml_paragraph_xml(formula))
            if note:
                set_run(p.add_run("　" + note), size=10.5)
            i += 1
            continue
        # 普通段落
        p = para(doc, "", size=12, spacing=1.5, space_after=6)
        render_inline(p, line.strip(), 12)
        i += 1

    # ---------------- 参考文献 ----------------
    doc.add_page_break()
    para(doc, "参考文献", size=14, ea=HEI, align=WD_ALIGN_PARAGRAPH.CENTER,
         spacing=1.5, space_before=6, space_after=10)
    for ref in references:
        if ref.startswith("注："):
            para(doc, ref, size=10.5, spacing=1.3, space_after=4)
            continue
        p = para(doc, "", size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT,
                 spacing=1.3, space_after=4)
        render_inline(p, ref, 10.5)

    path = OUT / DOCX_NAME
    tmp = OUT / "_paper_new.docx"
    doc.save(str(tmp))
    mo.fix_docx_namespaces(tmp)       # 关键：把 xmlns:m 搬到文档根元素，否则 Word 会丢弃公式
    try:
        os.replace(tmp, path)
        final = path
    except PermissionError:
        # 目标文件正被 Word 打开时无法覆盖：另存为修正版，避免白跑一遍
        alt = OUT / DOCX_NAME.replace(".docx", "·修正版.docx")
        os.replace(tmp, alt)
        final = alt
        print("[warn] 目标文件被占用（可能正在 Word/WPS 中打开），已另存为修正版")
    print(f"[OK] DOCX: {final}  ({final.stat().st_size/1024:.0f} KB)")
    print(f"     图片 {fig_no} 张，表格 {tbl_no} 个")
    return final


if __name__ == "__main__":
    build()
